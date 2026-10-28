"""Rendering: one-page ATS-friendly resume (PDF + DOCX) and a cover note."""

from __future__ import annotations

import copy
import re
import unicodedata
from pathlib import Path

from . import profile as pr
from . import skills as sk
from .jobs import Job

SKILL_GROUPS = [("Languages", ("language",)), ("Frameworks & Libraries", ("frontend", "backend")),
                ("Databases", ("database",)), ("Cloud & DevOps", ("cloud",)), ("Data & ML", ("data",)),
                ("Tools", ("tools",)), ("Concepts", ("concept",)), ("Other", ("other", "soft"))]
_ASCII = {"‘": "'", "’": "'", "“": '"', "”": '"', "–": "-", "—": "-",
          "…": "...", "•": "-", " ": " ", "→": "->", "™": "(TM)"}


def _latin1(text: str) -> str:
    text = "".join(_ASCII.get(c, c) for c in str(text))
    out = []
    for c in text:
        try:
            c.encode("latin-1")
            out.append(c)
        except UnicodeEncodeError:
            out.append(unicodedata.normalize("NFKD", c).encode("latin-1", "ignore").decode("latin-1"))
    return "".join(out)


def grouped_skills(skills: list[str]) -> list[tuple[str, list[str]]]:
    groups = []
    for label, cats in SKILL_GROUPS:
        items = [s for s in skills if sk.category(s) in cats]
        if items:
            groups.append((label, items))
    return groups


def resume_doc(profile: dict, tailored: dict) -> dict:
    """Merge profile + tailored picks into the exact content to print."""
    by_id = {e["id"]: e for sec in pr.SECTIONS for e in profile[sec]}
    exp = [{**by_id[t["id"]], "bullets": t.get("bullets") or by_id[t["id"]]["bullets"]}
           for t in tailored.get("experience", []) if t.get("id") in by_id]
    proj = [{**by_id[t["id"]], "bullets": t.get("bullets") or by_id[t["id"]]["bullets"]}
            for t in tailored.get("projects", []) if t.get("id") in by_id]
    real_work = any(e["bullets"] for e in exp)
    return {
        "name": profile["name"], "contact": _contact(profile), "headline": tailored.get("headline", ""),
        "summary": tailored.get("summary", ""), "education": profile["education"],
        "skills": grouped_skills(tailored.get("skills") or profile["skills"]),
        "experience": exp, "projects": proj,
        "achievements": tailored.get("achievements") or [], "certifications": profile.get("certifications") or [],
        "order": (["summary", "education", "skills", "experience", "projects", "achievements"] if real_work
                  else ["summary", "education", "skills", "projects", "experience", "achievements"]),
    }


def _contact(p: dict) -> list[tuple[str, str]]:
    out = [(p.get("email", ""), f"mailto:{p.get('email', '')}"), (p.get("phone", ""), ""), (p.get("location", ""), "")]
    for _, url in (p.get("links") or {}).items():
        out.append((re.sub(r"^https?://(www\.)?", "", url).rstrip("/"), url))
    return [(t, l) for t, l in out if t]


def _dates(e: dict) -> str:
    return " - ".join(filter(None, [e.get("start", ""), e.get("end", "")]))


# ----------------------------------------------------------------- PDF --

def _build_pdf(doc: dict, fs: float):
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(True, margin=12)
    pdf.set_margins(15, 12, 15)
    pdf.add_page()
    W = pdf.w - pdf.l_margin - pdf.r_margin
    lh = fs * 0.47
    NL = dict(new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.set_font("Helvetica", "B", fs + 9)
    pdf.cell(W, (fs + 9) * 0.5, _latin1(doc["name"]), align="C", **NL)
    contact = _latin1(" | ".join(t for t, _ in doc["contact"]))
    size = fs - 1
    pdf.set_font("Helvetica", "", size)
    while pdf.get_string_width(contact) > W and size > fs - 2.5:
        size -= 0.25
        pdf.set_font("Helvetica", "", size)
    pdf.multi_cell(W, lh + 1, contact, align="C", **NL)
    if doc["headline"]:
        pdf.set_font("Helvetica", "I", fs)
        pdf.cell(W, lh + 1, _latin1(doc["headline"]), align="C", **NL)

    def heading(text):
        pdf.ln(1.6)
        pdf.set_font("Helvetica", "B", fs + 1)
        pdf.cell(W, lh + 1, text.upper(), **NL)
        y = pdf.get_y()
        pdf.set_draw_color(60, 60, 60)
        pdf.line(pdf.l_margin, y, pdf.l_margin + W, y)
        pdf.ln(1)

    def left_right(left_bold, left_rest, right):
        pdf.set_font("Helvetica", "", fs - 0.5)
        rw = pdf.get_string_width(_latin1(right)) + 1 if right else 0
        pdf.set_font("Helvetica", "B", fs)
        lw = pdf.get_string_width(_latin1(left_bold))
        pdf.cell(lw + 0.5, lh + 0.6, _latin1(left_bold))
        pdf.set_font("Helvetica", "I", fs - 0.5)
        rest = _latin1(left_rest)
        while rest and pdf.get_string_width(rest) > W - rw - lw - 2:
            rest = rest[:-2]
        pdf.cell(W - rw - lw - 0.5, lh + 0.6, rest)
        pdf.set_font("Helvetica", "", fs - 0.5)
        pdf.cell(rw, lh + 0.6, _latin1(right), align="R", **NL)

    def bullets(items):
        pdf.set_font("Helvetica", "", fs)
        for b in items:
            y = pdf.get_y()
            pdf.set_fill_color(40, 40, 40)
            pdf.ellipse(pdf.l_margin + 1.6, y + lh / 2 - 0.45, 1.0, 1.0, style="F")
            pdf.set_x(pdf.l_margin + 4.5)
            pdf.multi_cell(W - 4.5, lh, _latin1(b), **NL)

    for part in doc["order"]:
        if part == "summary" and doc["summary"]:
            heading("Summary")
            pdf.set_font("Helvetica", "", fs)
            pdf.multi_cell(W, lh, _latin1(doc["summary"]), **NL)
        elif part == "education" and doc["education"]:
            heading("Education")
            for e in doc["education"]:
                left_right(e.get("school", ""), "", _dates(e))
                line = ", ".join(filter(None, [" ".join(filter(None, [e.get("degree", ""), e.get("field", "")])),
                                               f"CGPA/Grade: {e['grade']}" if e.get("grade") else ""]))
                if line:
                    pdf.set_font("Helvetica", "", fs)
                    pdf.multi_cell(W, lh, _latin1(line), **NL)
                if e.get("coursework"):
                    pdf.set_font("Helvetica", "I", fs - 0.5)
                    pdf.multi_cell(W, lh, _latin1("Coursework: " + ", ".join(e["coursework"])), **NL)
                bullets(e.get("bullets") or [])
        elif part == "skills" and doc["skills"]:
            heading("Technical Skills")
            for label, items in doc["skills"]:
                pdf.set_font("Helvetica", "B", fs)
                lw = pdf.get_string_width(label + ": ") + 0.5
                pdf.cell(lw, lh, _latin1(label + ": "))
                pdf.set_font("Helvetica", "", fs)
                pdf.multi_cell(W - lw, lh, _latin1(", ".join(items)), **NL)
        elif part == "experience" and doc["experience"]:
            heading("Experience")
            for e in doc["experience"]:
                left_right(e.get("role", ""), f"  {e.get('company', '')}" + (f", {e['location']}" if e.get("location") else ""),
                           _dates(e))
                bullets(e["bullets"])
        elif part == "projects" and doc["projects"]:
            heading("Projects")
            for e in doc["projects"]:
                left_right(e["name"], ("  |  " + ", ".join(e.get("tech", []))) if e.get("tech") else "",
                           re.sub(r"^https?://(www\.)?", "", e.get("link", "")))
                bullets(e["bullets"])
        elif part == "achievements" and (doc["achievements"] or doc["certifications"]):
            heading("Achievements & Certifications")
            bullets(doc["achievements"] + doc["certifications"])
    return pdf


def _save(write, path: str | Path) -> None:
    """Write via a temp file and swap it in, retrying while another program (OneDrive sync, a PDF
    viewer) briefly holds the old file open."""
    import os
    import time
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.stem}.tmp{path.suffix}")
    write(str(tmp))
    for attempt in range(8):
        try:
            os.replace(tmp, path)
            return
        except PermissionError:
            if attempt == 7:
                raise PermissionError(f"{path.name} is open in another program (close it and retry)") from None
            time.sleep(0.4)


def render_pdf(doc: dict, path: str | Path) -> dict:
    """Fit to one page: shrink type a little, then trim the least important bullets."""
    doc = copy.deepcopy(doc)
    trimmed: list[str] = []
    for _ in range(14):
        for fs in (10.5, 10, 9.5):
            pdf = _build_pdf(doc, fs)
            if pdf.page_no() == 1:
                _save(pdf.output, path)
                return {"path": str(path), "font_size": fs, "trimmed": trimmed, "pages": 1}
        # drop a bullet from the entry that has the most (projects/experience lists are relevance-ordered)
        pool = [e for e in doc["projects"] + doc["experience"] if len(e["bullets"]) > 1]
        if pool:
            e = max(pool, key=lambda x: (len(x["bullets"]), doc["projects"].index(x) if x in doc["projects"] else -1))
            trimmed.append(f"bullet from {e.get('name') or e.get('role')}: \"{e['bullets'][-1][:60]}…\"")
            e["bullets"] = e["bullets"][:-1]
        elif len(doc["projects"]) > 1:
            trimmed.append(f"project {doc['projects'][-1]['name']}")
            doc["projects"] = doc["projects"][:-1]
        elif doc["achievements"]:
            trimmed.append("an achievement")
            doc["achievements"] = doc["achievements"][:-1]
        else:
            break
    pdf = _build_pdf(doc, 9.5)
    _save(pdf.output, path)
    return {"path": str(path), "font_size": 9.5, "trimmed": trimmed, "pages": pdf.page_no()}


# ---------------------------------------------------------------- DOCX --

def render_docx(doc: dict, path: str | Path) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Mm(15)
        s.top_margin = s.bottom_margin = Mm(12)
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(0)
    right = d.sections[0].page_width - d.sections[0].left_margin - d.sections[0].right_margin

    def para(text="", bold=False, italic=False, size=None, align=None, space_before=0):
        p = d.add_paragraph()
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
        if align:
            p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        return p

    def heading(text):
        p = para(text.upper(), bold=True, size=11.5, space_before=6)
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "444444")):
            bottom.set(qn(k), v)
        border.append(bottom)
        p._p.get_or_add_pPr().append(border)

    def left_right(bold, rest, right_text):
        p = d.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(right, WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(bold).bold = True
        if rest:
            p.add_run(rest).italic = True
        if right_text:
            p.add_run("\t" + right_text)

    def bullets(items):
        for b in items:
            p = d.add_paragraph(b, style="List Bullet")
            p.paragraph_format.space_after = Pt(0)

    para(doc["name"], bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("  |  ".join(t for t, _ in doc["contact"]), size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    if doc["headline"]:
        para(doc["headline"], italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for part in doc["order"]:
        if part == "summary" and doc["summary"]:
            heading("Summary")
            para(doc["summary"])
        elif part == "education" and doc["education"]:
            heading("Education")
            for e in doc["education"]:
                left_right(e.get("school", ""), "", _dates(e))
                line = ", ".join(filter(None, [" ".join(filter(None, [e.get("degree", ""), e.get("field", "")])),
                                               f"CGPA/Grade: {e['grade']}" if e.get("grade") else ""]))
                if line:
                    para(line)
                if e.get("coursework"):
                    para("Coursework: " + ", ".join(e["coursework"]), italic=True)
                bullets(e.get("bullets") or [])
        elif part == "skills" and doc["skills"]:
            heading("Technical Skills")
            for label, items in doc["skills"]:
                p = d.add_paragraph()
                p.add_run(label + ": ").bold = True
                p.add_run(", ".join(items))
        elif part == "experience" and doc["experience"]:
            heading("Experience")
            for e in doc["experience"]:
                left_right(e.get("role", ""), f"  {e.get('company', '')}", _dates(e))
                bullets(e["bullets"])
        elif part == "projects" and doc["projects"]:
            heading("Projects")
            for e in doc["projects"]:
                left_right(e["name"], ("  |  " + ", ".join(e.get("tech", []))) if e.get("tech") else "",
                           re.sub(r"^https?://(www\.)?", "", e.get("link", "")))
                bullets(e["bullets"])
        elif part == "achievements" and (doc["achievements"] or doc["certifications"]):
            heading("Achievements & Certifications")
            bullets(doc["achievements"] + doc["certifications"])
    _save(d.save, path)
    return str(path)


# ---------------------------------------------------------- cover note --

def cover_note(profile: dict, job: Job, fit: dict, tailored: dict | None = None) -> str:
    """A short, honest cover note built only from profile facts."""
    edu = profile["education"][0] if profile["education"] else {}
    degree = " ".join(filter(None, [edu.get("degree", ""), edu.get("field", "")]))
    intro = f"I'm {profile['name']}" + (f", a {degree} student at {edu['school']}" if degree and edu.get("school") else "")
    by_id = {e["id"]: e for e in profile["projects"]}
    picks = [by_id[t["id"]] for t in (tailored or {}).get("projects", []) if t["id"] in by_id] or profile["projects"]
    body = [f"{intro}, and I'd like to apply for the {job.title} role."]
    if picks:
        p = picks[0]
        first = p["bullets"][0].rstrip(".") if p["bullets"] else ""
        body.append(f"Most relevant to this role is {p['name']}"
                    + (f" ({', '.join(p.get('tech', [])[:3])})" if p.get("tech") else "")
                    + (f": {first[0].lower() + first[1:]}." if first else "."))
    if fit["matched"]:
        body.append(f"I've worked hands-on with {', '.join(fit['matched'][:4])}, which this role asks for.")
    if fit["missing"]:
        body.append(f"I haven't used {fit['missing'][0]} in a project yet, and I'm learning it now.")
    body.append(f"I'd love to bring the same care to {job.company}.")
    sign = " · ".join(filter(None, [profile.get("email", ""), profile.get("phone", "")]))
    return "\n".join([f"Dear {job.company} hiring team,", "", " ".join(body), "",
                      "Thank you for your time,", profile["name"], sign])
